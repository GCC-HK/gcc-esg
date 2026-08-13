#!/usr/bin/env python3
"""Builds the CBAM methodology .docx from the HTML source.

Source of truth: scripts/cbam-methodology-source.html (edit there, then run
this script). Output: "CBAM Calculator - Methodology and Verification.docx"
in the repo root. Requires python-docx (pip install --target <dir> python-docx).
"""
import re
import sys
import os
from html.parser import HTMLParser

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
for p in sys.path[:]:
    lib = os.path.join(p, 'pylib')
    if os.path.isdir(lib):
        sys.path.insert(0, lib)
# scratchpad fallback used during development
sys.path.insert(0, '/private/tmp/claude-501/-Users-michaelkeppe-claude-projects-gcc-esg/a8db4631-fe74-4387-a298-0ff544d79c0c/scratchpad/pylib')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

GREEN = RGBColor(0x1A, 0x5C, 0x38)
BLUE = RGBColor(0x00, 0x33, 0x66)
AMBER = RGBColor(0x8A, 0x5A, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)

ROOT = os.path.join(os.path.dirname(__file__), '..')
SRC = os.path.join(ROOT, 'scripts', 'cbam-methodology-source.html')
OUT = os.path.join(ROOT, 'CBAM Calculator - Methodology and Verification.docx')


class Extract(HTMLParser):
    """Flattens the restricted HTML vocabulary of the source document into a
    list of (kind, payload) blocks; inline b/strong/i markers survive as
    segment flags."""

    BLOCK_TAGS = {'h1', 'h2', 'h3', 'p', 'li', 'td', 'th'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []
        self.stack = []
        self.cur = None          # list of (text, bold, italic) segments
        self.bold = 0
        self.italic = 0
        self.table = None        # current table rows
        self.row = None
        self.cellkind = None
        self.divclass = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'div' and a.get('class') in ('formula', 'caveat', 'resolved', 'meta'):
            self.divclass = a['class']
        if tag in ('b', 'strong'):
            self.bold += 1
        elif tag == 'i':
            self.italic += 1
        elif tag == 'br':
            if self.cur is not None:
                self.cur.append(('\n', False, False))
        elif tag == 'table':
            self.table = []
        elif tag == 'tr':
            self.row = []
        elif tag in ('td', 'th'):
            self.cur = []
            self.cellkind = tag
        elif tag in ('h1', 'h2', 'h3', 'li'):
            self.cur = []
            self.stack.append(tag)
        elif tag == 'p':
            self.cur = []
            self.stack.append('p-' + (self.divclass or (dict(attrs).get('class') or '')))

    def handle_endtag(self, tag):
        if tag in ('b', 'strong'):
            self.bold = max(0, self.bold - 1)
        elif tag == 'i':
            self.italic = max(0, self.italic - 1)
        elif tag in ('td', 'th'):
            self.row.append((self.cellkind, self.cur or []))
            self.cur = None
        elif tag == 'tr':
            if self.row:
                self.table.append(self.row)
            self.row = None
        elif tag == 'table':
            self.blocks.append(('table', self.table))
            self.table = None
        elif tag in ('h1', 'h2', 'h3', 'li'):
            kind = self.stack.pop()
            self.blocks.append((kind, self.cur or []))
            self.cur = None
        elif tag == 'p':
            kind = self.stack.pop()
            self.blocks.append((kind, self.cur or []))
            self.cur = None
        elif tag == 'div':
            self.divclass = None

    def handle_data(self, data):
        if self.cur is not None and data:
            text = re.sub(r'\s+', ' ', data)
            if text:
                self.cur.append((text, self.bold > 0, self.italic > 0))


def add_runs(par, segments, size=None, color=None, mono=False):
    for text, bold, italic in segments:
        for i, piece in enumerate(text.split('\n')):
            if i:
                par.add_run().add_break()
            if not piece:
                continue
            r = par.add_run(piece)
            r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
            if mono:
                r.font.name = 'Consolas'


def main():
    html = open(SRC, encoding='utf-8').read()
    body = html[html.index('<body>'):]
    ex = Extract()
    ex.feed(body)

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.2)
        section.left_margin = section.right_margin = Cm(2.2)
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    for kind, payload in ex.blocks:
        if kind == 'h1':
            p = doc.add_heading(level=0)
            add_runs(p, payload, size=17, color=GREEN)
        elif kind == 'h2':
            p = doc.add_heading(level=1)
            add_runs(p, payload, size=13, color=GREEN)
        elif kind == 'h3':
            p = doc.add_heading(level=2)
            add_runs(p, payload, size=11.5, color=BLUE)
        elif kind == 'li':
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, payload)
        elif kind == 'table':
            rows, cols = len(payload), max(len(r) for r in payload)
            t = doc.add_table(rows=rows, cols=cols)
            t.style = 'Table Grid'
            t.autofit = True
            for ri, row in enumerate(payload):
                for ci, (cellkind, segs) in enumerate(row):
                    cell = t.rows[ri].cells[ci]
                    cp = cell.paragraphs[0]
                    if cellkind == 'th':
                        add_runs(cp, [(s[0], True, s[2]) for s in segs], size=9.5, color=GREEN)
                    else:
                        add_runs(cp, segs, size=9.5)
            doc.add_paragraph()
        elif kind.startswith('p'):
            cls = kind[2:]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if cls == 'formula':
                add_runs(p, payload, size=9.5, mono=True)
                p.paragraph_format.left_indent = Cm(0.5)
            elif cls in ('caveat', 'resolved'):
                add_runs(p, payload, size=10, color=AMBER if cls == 'caveat' else GREEN)
                p.paragraph_format.left_indent = Cm(0.5)
            elif cls == 'meta':
                add_runs(p, payload, size=9.5, color=GREY)
            else:
                add_runs(p, payload)

    doc.save(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
