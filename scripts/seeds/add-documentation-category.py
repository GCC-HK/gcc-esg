#!/usr/bin/env python3
"""Owner decision 2026-09-19: add the matchmaking category "Technical
documentation and product information" with Impala Services and Pergamon Labs
as providers. The private workbook is the single source of truth for
matchmaking data, so the rows are added there (backup written first), then
scripts/build-matchmaking-data.py regenerates the public dataset.

Idempotent: python3 scripts/seeds/add-documentation-category.py
Offerings reuse the companies' existing verified offering texts (both already
name technical documentation); nothing is invented.
"""
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent.parent
WB_PATH = ROOT / "input" / "GCC-Members-ESG-Mapping.xlsx"

CATEGORY = "Technical documentation and product information"
EXAMPLE = "We need help preparing technical files, manuals or product information for EU requirements."

PROVIDERS = [
    {
        "Company": "Impala Services Ltd.",
        "Offering": "Supports multilingual product communication, technical documentation and digital compliance content relevant to DPP and sustainable-product information workflows.",
        "Keywords": "technical documentation; product information; multilingual content; manuals; DPP content",
        "Website": "https://www.impala-tech.com/",
    },
    {
        "Company": "Pergamon Labs",
        "Offering": "Provides AI-powered structured documentation and disclosure workflows relevant to EU compliance, packaging, manuals and Digital Product Passports.",
        "Keywords": "technical documentation; manuals; structured content; disclosure; DPP",
        "Website": "https://www.pergamon-labs.com/",
    },
]


def main():
    if not WB_PATH.exists():
        sys.exit(f"Workbook not found: {WB_PATH}")
    wb = openpyxl.load_workbook(WB_PATH)

    tax = wb["Matchmaking Taxonomy"]
    tax_cats = {r[0] for r in tax.iter_rows(min_row=2, values_only=True) if r[0]}
    if CATEGORY not in tax_cats:
        tax.append([CATEGORY, EXAMPLE])
        print(f"Taxonomy: added '{CATEGORY}'")
    else:
        print(f"Taxonomy: '{CATEGORY}' already present")

    prov = wb["Providers by Category"]
    header = [c.value for c in prov[1]]
    existing = {(r[header.index("Category")], r[header.index("Company")])
                for r in prov.iter_rows(min_row=2, values_only=True) if r[0]}
    added = 0
    for p in PROVIDERS:
        if (CATEGORY, p["Company"]) in existing:
            print(f"Providers: {p['Company']} already listed under the category")
            continue
        row = {"Category": CATEGORY, "Company": p["Company"], "ESG Provider": "Yes",
               "Offering": p["Offering"], "Keywords": p["Keywords"], "Website": p["Website"],
               "Confidence": "High", "Partner Potential": "High",
               "Partner Role": "Platform and implementation partner",
               "Source": p["Website"]}
        prov.append([row.get(h, "") for h in header])
        added += 1
        print(f"Providers: added {p['Company']} under '{CATEGORY}'")

    if added or CATEGORY not in tax_cats:
        backup = WB_PATH.with_name(f"GCC-Members-ESG-Mapping.backup-{date.today().isoformat()}.xlsx")
        if not backup.exists():
            shutil.copy2(WB_PATH, backup)
            print(f"Backup written: {backup.name}")
        wb.save(WB_PATH)
        print("Workbook saved. Now run: python3 scripts/build-matchmaking-data.py")
    else:
        print("No changes needed.")


if __name__ == "__main__":
    main()
